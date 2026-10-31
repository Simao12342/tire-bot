import os
import re
import sys
import io
import json
import html
import asyncio
import base64
import difflib
import threading
import time as _time
from datetime import datetime, timedelta, time as dtime
from http.server import HTTPServer, BaseHTTPRequestHandler
from zoneinfo import ZoneInfo

import pandas as pd
import gspread
from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    ReplyKeyboardMarkup,
)
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ConversationHandler,
    filters,
    ContextTypes,
)

# ----------------- ФОНОВИЙ ВЕБ-СЕРВЕР ДЛЯ RENDER -----------------
class HealthCheckHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"OK")

    def do_HEAD(self):
        self.send_response(200)
        self.end_headers()

    def log_message(self, *args):
        pass


def run_health_check_server():
    port = int(os.environ.get("PORT", 8080))
    HTTPServer(("0.0.0.0", port), HealthCheckHandler).serve_forever()


threading.Thread(target=run_health_check_server, daemon=True).start()

# ----------------- НАЛАШТУВАННЯ -----------------
# Токен береться зі змінної середовища TELEGRAM_BOT_TOKEN (Render -> Environment).
# Якщо змінної немає — встав токен між лапками нижче.
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN") or "ВСТАВ_СЮДИ_ТОКЕН"
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")      # потрібен ТІЛЬКИ для пошуку по фото цінника
SPREADSHEET_NAME = os.environ.get("SPREADSHEET_NAME", "шини 2026")

# Склад для звірки: за замовчуванням аркуш «Склад» у цій самій Google Таблиці
# (колонки: Артикул | Назва | Кількість). Можна змінити назву аркуша змінною SKLAD_SHEET
# або вказати окрему таблицю змінною SKLAD_SPREADSHEET (назва, посилання або ID).
SKLAD_SPREADSHEET = os.environ.get("SKLAD_SPREADSHEET", "").strip()
SKLAD_SHEET = os.environ.get("SKLAD_SHEET", "").strip() or "Склад"

# Необов'язково: чати для щоденного звіту й сповіщень про мінімум (через кому). Свій ID: команда /myid
REPORT_CHAT_IDS = [int(x) for x in re.findall(r"-?\d+", os.environ.get("REPORT_CHAT_IDS", ""))]
REPORT_TIME = os.environ.get("REPORT_TIME", "08:30")        # щоденний звіт, київський час
LOW_STOCK = int(os.environ.get("LOW_STOCK", "2"))           # мінімальний залишок по артикулу (разом)

FIRST_DATA_ROW = 4      # дані в таблиці починаються з 4-го рядка
MAX_CARDS = 6           # скільки карток показувати одразу (інакше — список кнопок)
MAX_LIST = 40           # скільки артикулів у списку кнопок
MAX_QTY = 999           # більше — схоже, що замість кількості вписали артикул
SANE_QTY = 5000         # у звітах ігноруємо записи історії з більшими значеннями
CACHE_TTL = 45          # секунд, кеш таблиці
SKLAD_TTL = 120
HISTORY_TTL = 60
UNDO_WINDOW = 15 * 60   # секунд, скільки діє кнопка "Відмінити"
KYIV = ZoneInfo("Europe/Kyiv")

openai_client = None
if OPENAI_API_KEY:
    try:
        import openai
        openai_client = openai.OpenAI(api_key=OPENAI_API_KEY)
    except Exception as e:
        print(f"⚠️ OpenAI недоступний, фото-пошук вимкнено: {e}")

SPECIAL_LOCS = {          # короткі ключі для кнопок -> назви в таблиці
    "zal": "вітрина зал",
    "vul": "вітрина вулиця",
    "fas": "палети фасад",
}


class StaleRow(Exception):
    """Рядок у таблиці змінився (хтось видалив/перемістив рядки)."""


# ----------------- GOOGLE ТАБЛИЦЯ + КЕШ -----------------
_lock = threading.RLock()
_gc = None
_spreadsheet = None
_cache = {"df": None, "ts": 0.0}
_sklad_cache = {"data": None, "ts": 0.0}
_hist_cache = {"rows": None, "ts": 0.0}


def get_client():
    global _gc
    with _lock:
        if _gc is None:
            creds_json = os.environ.get("GOOGLE_CREDENTIALS_JSON")
            if creds_json:
                _gc = gspread.service_account_from_dict(json.loads(creds_json))
            else:
                _gc = gspread.service_account(filename="credentials.json")
        return _gc


def service_email():
    try:
        creds_json = os.environ.get("GOOGLE_CREDENTIALS_JSON")
        if creds_json:
            return json.loads(creds_json).get("client_email", "")
        with open("credentials.json", encoding="utf-8") as f:
            return json.load(f).get("client_email", "")
    except Exception:
        return ""


def get_spreadsheet():
    global _spreadsheet
    with _lock:
        if _spreadsheet is None:
            _spreadsheet = get_client().open(SPREADSHEET_NAME)
        return _spreadsheet


def get_worksheet():
    return get_spreadsheet().sheet1  # перший аркуш ("Шини")


def to_int(v):
    s = str(v).strip().replace(",", ".")
    try:
        return int(float(s))
    except Exception:
        return 0


def load_df(force=False):
    """Таблиця «Шини» як DataFrame. Кешується на CACHE_TTL секунд."""
    with _lock:
        now = _time.time()
        if not force and _cache["df"] is not None and now - _cache["ts"] < CACHE_TTL:
            return _cache["df"].copy()
    data = get_worksheet().get_all_values()
    rows = []
    for i, r in enumerate(data[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        r = (list(r) + [""] * 5)[:5]
        rows.append([i] + [str(x).strip() for x in r])
    df = pd.DataFrame(rows, columns=["row", "pallet", "code", "size", "name", "qty"])
    df["code"] = df["code"].str.replace(r"\.0$", "", regex=True)
    df = df[df["code"] != ""].copy()
    df["q"] = df["qty"].apply(to_int)
    with _lock:
        _cache["df"] = df
        _cache["ts"] = _time.time()
    return df.copy()


def invalidate_cache():
    with _lock:
        _cache["ts"] = 0.0
        _hist_cache["ts"] = 0.0


def cache_set_qty(row, qty):
    with _lock:
        df = _cache["df"]
        if df is not None:
            df.loc[df["row"] == row, "qty"] = str(qty)
            df.loc[df["row"] == row, "q"] = int(qty)


def cached_code_of_row(row):
    with _lock:
        df = _cache["df"]
        if df is None:
            return None
        m = df[df["row"] == row]
        return None if m.empty else str(m.iloc[0]["code"])


# ----------------- СКЛАД (необов'язково) -----------------
_sklad_meta = {"rows": 0, "dups": []}


def load_sklad(force=False):
    """dict: артикул -> {'name', 'qty'}; None, якщо аркуша зі складом немає. Дублі артикулів підсумовуються."""
    with _lock:
        if not force and _time.time() - _sklad_cache["ts"] < SKLAD_TTL and _sklad_cache["ts"] > 0:
            return _sklad_cache["data"]
    try:
        if SKLAD_SPREADSHEET:
            gc = get_client()
            m = re.search(r"/d/([A-Za-z0-9_-]+)", SKLAD_SPREADSHEET)
            if m:
                sh = gc.open_by_key(m.group(1))
            elif re.fullmatch(r"[A-Za-z0-9_-]{30,}", SKLAD_SPREADSHEET):
                sh = gc.open_by_key(SKLAD_SPREADSHEET)
            else:
                sh = gc.open(SKLAD_SPREADSHEET)
            ws = sh.worksheet(SKLAD_SHEET) if os.environ.get("SKLAD_SHEET", "").strip() else sh.sheet1
        else:
            ws = get_spreadsheet().worksheet(SKLAD_SHEET)
    except gspread.exceptions.WorksheetNotFound:
        with _lock:
            _sklad_cache["data"], _sklad_cache["ts"] = None, _time.time()
        return None
    result, seen, dups, n_rows = {}, {}, [], 0
    for r in ws.get_all_values():
        cells = [str(c).strip() for c in r]
        code_idx = next((i for i, s_ in enumerate(cells) if re.fullmatch(r"\d{6,9}(\.0)?", s_)), None)
        if code_idx is None:
            continue
        code = cells[code_idx].split(".")[0]
        qty = None
        for i in range(len(cells) - 1, -1, -1):
            if i != code_idx and re.fullmatch(r"-?\d+([.,]0+)?", cells[i]):
                qty = to_int(cells[i])
                break
        texts = [s_ for i, s_ in enumerate(cells) if i != code_idx and s_ and not re.fullmatch(r"-?[\d.,]+", s_)]
        name = max(texts, key=len) if texts else ""
        n_rows += 1
        e = result.setdefault(code, {"name": name, "qty": 0})
        e["qty"] += qty or 0
        if not e["name"]:
            e["name"] = name
        seen[code] = seen.get(code, 0) + 1
        if seen[code] == 2:
            dups.append(code)
    with _lock:
        _sklad_cache["data"], _sklad_cache["ts"] = result, _time.time()
        _sklad_meta["rows"], _sklad_meta["dups"] = n_rows, dups
    return result


def load_sklad_safe(force=False):
    try:
        data = load_sklad(force)
        return data if data else None
    except Exception as e:
        print(f"⚠️ Не вдалося прочитати склад: {e}")
        return None


# ----------------- ІСТОРІЯ -----------------
def now_str():
    return datetime.now(KYIV).strftime("%Y-%m-%d %H:%M:%S")


def user_label(user):
    return f"@{user.username}" if user.username else f"{user.first_name} (ID: {user.id})"


ACTION_TEXT = {
    "inc": "Додано (+1)",
    "dec": "Зменшено (-1)",
    "add": "Додано нову шину",
    "merge": "Додано кількість до наявної",
    "undo": "Скасовано зміну",
    "inv": "Інвентаризація",
    "set": "Встановлено точну кількість",
}


def log_to_history(user, code, action, old_qty, new_qty):
    try:
        history = get_spreadsheet().worksheet("Історія")
        history.append_row([now_str(), user_label(user), str(code), ACTION_TEXT.get(action, ACTION_TEXT["set"]), old_qty, new_qty])
        _hist_cache["ts"] = 0.0
    except Exception as e:
        print(f"⚠️ Помилка запису в історію: {e}")


def read_history():
    """Список записів історії: dict(ts, user, code, action, old, new)."""
    with _lock:
        if _hist_cache["rows"] is not None and _time.time() - _hist_cache["ts"] < HISTORY_TTL:
            return _hist_cache["rows"]
    try:
        data = get_spreadsheet().worksheet("Історія").get_all_values()
    except Exception as e:
        print(f"⚠️ Не вдалося прочитати історію: {e}")
        return []
    out = []
    for r in data[1:]:
        r = (list(r) + [""] * 6)[:6]
        try:
            ts = datetime.strptime(r[0].strip(), "%Y-%m-%d %H:%M:%S")
        except Exception:
            continue
        out.append({"ts": ts, "user": r[1].strip(), "code": r[2].strip(), "action": r[3].strip(),
                    "old": to_int(r[4]), "new": to_int(r[5])})
    with _lock:
        _hist_cache["rows"] = out
        _hist_cache["ts"] = _time.time()
    return out


# ----------------- ЗАПИС У ТАБЛИЦЮ -----------------
def change_qty(user, row, action, value=None, log=True):
    ws = get_worksheet()
    vals = (ws.row_values(row) + [""] * 5)[:5]
    code = str(vals[1]).strip().replace(".0", "")
    expected = cached_code_of_row(row)
    if not code or (expected is not None and expected != code):
        invalidate_cache()
        raise StaleRow()
    old = to_int(vals[4])
    if action == "inc":
        new = old + 1
    elif action == "dec":
        new = max(0, old - 1)
    else:
        new = int(value)
    ws.update_cell(row, 5, new)
    cache_set_qty(row, new)
    if log:
        log_to_history(user, code, action, old, new)
    return code, old, new


def apply_qty_batch(user, changes, action="inv"):
    """changes: список (row, code, old, new). Один запит на запис і один на історію."""
    if not changes:
        return 0
    ws = get_worksheet()
    ws.batch_update(
        [{"range": f"E{row}", "values": [[new]]} for row, _code, _old, new in changes],
        value_input_option="USER_ENTERED",
    )
    for row, _code, _old, new in changes:
        cache_set_qty(row, new)
    try:
        history = get_spreadsheet().worksheet("Історія")
        ts, who = now_str(), user_label(user)
        history.append_rows(
            [[ts, who, str(code), ACTION_TEXT[action], old, new] for _row, code, old, new in changes],
            value_input_option="USER_ENTERED",
        )
        _hist_cache["ts"] = 0.0
    except Exception as e:
        print(f"⚠️ Помилка запису в історію: {e}")
    return len(changes)


def add_tire(user, pallet, code, size, name, qty):
    ws = get_worksheet()
    next_row = max(len(ws.col_values(2)) + 1, FIRST_DATA_ROW)  # перший вільний рядок після останнього артикула
    if ws.row_count < next_row:
        ws.add_rows(next_row - ws.row_count)
    ws.update(
        range_name=f"A{next_row}:E{next_row}",
        values=[[pallet, code, size, name, qty]],
        value_input_option="USER_ENTERED",
    )
    saved = (ws.row_values(next_row) + [""] * 5)[:5]
    if str(saved[1]).strip() != str(code).strip():
        raise RuntimeError(f"запис не з'явився в рядку {next_row}")
    invalidate_cache()
    log_to_history(user, code, "add", 0, qty)
    return next_row


def merge_qty(user, row, extra):
    ws = get_worksheet()
    vals = (ws.row_values(row) + [""] * 5)[:5]
    code = str(vals[1]).strip().replace(".0", "")
    old = to_int(vals[4])
    new = old + int(extra)
    ws.update_cell(row, 5, new)
    cache_set_qty(row, new)
    log_to_history(user, code, "merge", old, new)
    return code, old, new


# ----------------- ДОПОМІЖНЕ -----------------
def short_name(name):
    return re.sub(r"^\s*Автошина\s+", "", str(name)).strip()


def loc_label(pallet):
    p = str(pallet).strip()
    return f"палета {p}" if p.isdigit() else (p or "без локації")


def loc_sort_key(p):
    p = str(p).strip()
    return (0, int(p), "") if p.isdigit() else (1, 0, p)


def friendly_error(e):
    print(f"⚠️ Помилка: {type(e).__name__}: {e}")
    text = str(e)
    if isinstance(e, StaleRow):
        return "🔄 Таблиця змінилась, поки ти працював. Знайди шину ще раз і повтори дію."
    if "429" in text or "Quota" in text or "quota" in text:
        return "⏳ Google просить трохи зачекати. Спробуй ще раз за хвилину."
    if "403" in text or "PERMISSION" in text.upper():
        return "🔒 Немає доступу до таблиці. Перевір, що їй надано доступ для сервісного акаунта."
    if "404" in text or "SpreadsheetNotFound" in type(e).__name__:
        return "❓ Не знайшов таблицю. Перевір назву в налаштуваннях бота."
    return "⚠️ Щось пішло не так. Спробуй ще раз; якщо повториться, напиши адміністратору."


def chunk_text(text, limit=3900):
    parts, cur = [], ""
    for line in text.split("\n"):
        if len(cur) + len(line) + 1 > limit:
            parts.append(cur)
            cur = ""
        cur += line + "\n"
    if cur.strip():
        parts.append(cur)
    return parts


def get_tire_photo(name):
    photos_dir = "photo"
    if not os.path.exists(photos_dir):
        return None
    name_lower = str(name).lower()
    for filename in os.listdir(photos_dir):
        if filename.lower().endswith((".jpg", ".jpeg", ".png", ".webp", ".avif")):
            keyword = os.path.splitext(filename)[0].lower()
            if len(keyword) >= 3 and keyword in name_lower:
                return os.path.join(photos_dir, filename)
    return None


# ----------------- ПОШУК -----------------
BRAND_ALIASES = {
    "петлас": "petlas", "хенкук": "hankook", "ханкук": "hankook", "хенкок": "hankook",
    "тріангл": "triangle", "трайангл": "triangle", "триангл": "triangle",
    "маршал": "marshal", "росава": "rosava", "аосен": "aosen", "хайфлай": "hifly", "хіфлай": "hifly",
    "преміорі": "premiorri", "премиорри": "premiorri", "премьорі": "premiorri", "естрада": "estrada",
    "горизонт": "horizon", "хорайзон": "horizon", "павертрак": "powertrac", "пауертрак": "powertrac",
    "матадор": "matador", "мішлен": "michelin", "бріджстоун": "bridgestone", "нокіан": "nokian",
}
WORD_ALIASES = {"зим": "зима", "літ": "літо", "всесез": "всесезон"}
SIZE_RE = re.compile(r"(?<!\d)(\d{3})\s*[/\-\s.xх]?\s*(\d{2})\s*[/\-\s]?\s*[rрzZ]?\s*(\d{2})(?!\d)", re.I)


def norm(s):
    """Лишає тільки літери й цифри: '205/55 R16' -> '20555r16'."""
    return re.sub(r"[^0-9a-zа-яіїєґ]", "", str(s).lower())


def alias(t):
    for k, v in BRAND_ALIASES.items():
        if t.startswith(k):
            return v
    for k, v in WORD_ALIASES.items():
        if t.startswith(k):
            return v
    return t


def parse_query(q):
    q = str(q).strip().lower()
    size_key = None
    m = SIZE_RE.search(q)
    if m:
        size_key = f"{m.group(1)}{m.group(2)}r{m.group(3)}"
        q = (q[:m.start()] + " " + q[m.end():]).strip()
    tokens = []
    for t in re.split(r"\s+", q):
        t = t.strip(",.;:()[]\"'")
        if not t:
            continue
        if re.search(r"\d", t):
            t = t.replace("р", "r")
        tokens.append(alias(t))
    return size_key, tokens


def _match(df, size_key, tokens):
    hay = (df["code"] + " " + df["size"] + " " + df["name"]).str.lower()
    hay_c = hay.apply(norm)
    mask = pd.Series(True, index=df.index)
    if size_key:
        mask &= hay_c.str.contains(size_key, regex=False)
    for t in tokens:
        mask &= hay.str.contains(t, regex=False) | hay_c.str.contains(norm(t) or t, regex=False)
    return df[mask]


def fuzzy_fix(df, tokens):
    vocab = set()
    for s in df["name"].tolist():
        for w in re.findall(r"[A-Za-z][A-Za-z0-9+]{3,}", s):
            vocab.add(w.lower())
    vocab_l = sorted(vocab)
    fixed, changes = [], []
    for t in tokens:
        if len(t) < 4 or any(ch.isdigit() for ch in t) or any(t in v for v in vocab_l):
            fixed.append(t)
            continue
        m = difflib.get_close_matches(t, vocab_l, n=1, cutoff=0.72)
        if m:
            fixed.append(m[0])
            changes.append((t, m[0]))
        else:
            fixed.append(t)
    return fixed, changes


def search_codes(df, query):
    """Повертає (список артикулів, примітка)."""
    if df.empty:
        return [], ""
    size_key, tokens = parse_query(query)
    if not size_key and not tokens:
        return [], ""
    found = _match(df, size_key, tokens)
    note = ""
    if found.empty and tokens:
        fixed, changes = fuzzy_fix(df, tokens)
        if changes:
            found = _match(df, size_key, fixed)
            if not found.empty:
                note = "🔎 Шукаю " + ", ".join(f"«{b}» замість «{a}»" for a, b in changes)
    if found.empty:
        digits = re.sub(r"\D", "", str(query))
        if len(digits) >= 4:
            found = df[df["code"].str.contains(digits, regex=False)]
    codes = list(dict.fromkeys(found["code"].tolist()))
    q = str(query).strip()
    if q in codes:  # точний збіг артикула — першим
        codes.remove(q)
        codes.insert(0, q)
    return codes, note


# ----------------- КАРТКА ТОВАРУ -----------------
def code_rows(df, code):
    rows = df[df["code"] == code].copy()
    rows["_k"] = rows["pallet"].apply(loc_sort_key)
    return rows.sort_values(by=["_k", "row"]).drop(columns=["_k"])


def code_total(df, code):
    return int(df[df["code"] == code]["q"].sum())


def build_card(code, df, sk=None, undo_row=None):
    rows = code_rows(df, code)
    if rows.empty:
        return f"❓ Артикул {code} не знайдено в таблиці.", None
    first = rows.iloc[0]
    names = {norm(short_name(n)) for n in rows["name"]}
    sizes = {norm(s) for s in rows["size"] if s}
    total = int(rows["q"].sum())
    labels = [loc_label(p) for p in rows["pallet"]]
    dup = {l for l in labels if labels.count(l) > 1}

    lines, kb = [], []
    for (_, r), lab in zip(rows.iterrows(), labels):
        row = int(r["row"])
        shown = f"{lab} (р.{row})" if lab in dup else lab
        q = int(r["q"])
        lines.append(f"• {shown} — {q} шт.")
        kb.append([
            InlineKeyboardButton("➖", callback_data=f"dec_{row}"),
            InlineKeyboardButton(f"{shown}: {q} шт.", callback_data=f"set_{row}"),
            InlineKeyboardButton("➕", callback_data=f"inc_{row}"),
        ])

    text = (
        f"🛞 {short_name(first['name'])}\n"
        f"🔢 Артикул: {code}   📏 {first['size'] or '—'}\n\n"
        "📍 Де лежить:\n" + "\n".join(lines) + "\n\n"
        f"📦 Разом у таблиці: {total} шт."
    )
    if sk is not None:
        if code in sk:
            diff = total - sk[code]["qty"]
            text += f"\n🏬 Склад: {sk[code]['qty']} шт. — " + ("✅ збігається" if diff == 0 else f"⚠️ у таблиці {diff:+d}")
        else:
            text += "\n🏬 Склад: такого артикула немає в базі"
    if total == 0:
        text += "\n⚠️ Немає в наявності"
    elif total <= LOW_STOCK:
        text += f"\n🔻 Мінімальний залишок (≤ {LOW_STOCK} шт. разом)"
    if len(names) > 1 or len(sizes) > 1:
        text += "\n⚠️ У рядках різні назви або розміри: варто перевірити"

    footer = []
    if undo_row:
        footer.append(InlineKeyboardButton("↩️ Відмінити", callback_data=f"undo_{undo_row}"))
    footer.append(InlineKeyboardButton("🔄 Оновити", callback_data=f"rf_{code}"))
    kb.append(footer)
    return text, InlineKeyboardMarkup(kb)


def undo_row_for(context, df, code):
    last = context.application.bot_data.get("last", {})
    rows = set(df[df["code"] == code]["row"].astype(int).tolist())
    now = _time.time()
    cands = [(v["ts"], r) for r, v in last.items() if r in rows and now - v["ts"] < UNDO_WINDOW]
    return max(cands)[1] if cands else None


async def card_payload(context, code, force=False):
    df = await asyncio.to_thread(load_df, force)
    sk = await asyncio.to_thread(load_sklad_safe)
    return build_card(code, df, sk, undo_row_for(context, df, code))


async def send_card(chat_send, context, code, force=False):
    text, kb = await card_payload(context, code, force)
    return await chat_send(text, reply_markup=kb)


def list_keyboard(df, codes):
    kb = []
    for c in codes[:MAX_LIST]:
        rows = df[df["code"] == c]
        if rows.empty:
            continue
        label = f"{c} · {rows.iloc[0]['size']} · {short_name(rows.iloc[0]['name'])}"
        kb.append([InlineKeyboardButton(f"{label[:52]} · {int(rows['q'].sum())} шт.", callback_data=f"card_{c}")])
    return InlineKeyboardMarkup(kb)


# ----------------- ЛОКАЦІЇ -----------------
def resolve_loc_key(text):
    t = str(text).strip().lower()
    if not t:
        return None
    if t.isdigit():
        return t
    m = re.fullmatch(r"(?:палета|палет|пал\.?|№)\s*(\d+)", t)
    if m:
        return m.group(1)
    for k, v in SPECIAL_LOCS.items():
        if t == k or t in v:
            return k
    return None


def loc_name_of_key(key):
    return SPECIAL_LOCS.get(key) or loc_label(key)


def location_rows(df, key):
    if key.isdigit():
        return df[df["pallet"].str.strip() == key].sort_values("row")
    name = SPECIAL_LOCS.get(key, key).lower()
    return df[df["pallet"].str.strip().str.lower() == name].sort_values("row")


def palette_keyboard(prefix):
    nums = [InlineKeyboardButton(str(n), callback_data=f"{prefix}_{n}") for n in range(1, 45)]
    kb = [nums[i:i + 6] for i in range(0, len(nums), 6)]
    kb.append([
        InlineKeyboardButton("🪟 Вітрина зал", callback_data=f"{prefix}_zal"),
        InlineKeyboardButton("🏬 Вулиця", callback_data=f"{prefix}_vul"),
        InlineKeyboardButton("🏗 Фасад", callback_data=f"{prefix}_fas"),
    ])
    return InlineKeyboardMarkup(kb)


def build_location_view(df, key):
    rows = location_rows(df, key)
    label = loc_name_of_key(key)
    if rows.empty:
        return f"📦 {label}: тут нічого не записано.", None, rows
    total = int(rows["q"].sum())
    lines = [f"📦 {label} — {len(rows)} позицій, {total} шт.\n"]
    for _, r in rows.iterrows():
        lines.append(f"• {r['code']} · {r['size']} · {short_name(r['name'])} — {int(r['q'])} шт.")
    kb = []
    if len(rows) <= 24:
        btns = [InlineKeyboardButton(f"{r['code']} · {int(r['q'])}", callback_data=f"card_{r['code']}") for _, r in rows.iterrows()]
        kb += [btns[i:i + 2] for i in range(0, len(btns), 2)]
    kb.append([InlineKeyboardButton("🧮 Інвентаризація", callback_data=f"invs_{key}")])
    kb.append([InlineKeyboardButton("⬅️ Усі палети", callback_data="pals")])
    return "\n".join(lines), InlineKeyboardMarkup(kb), rows


# ----------------- СТАТИСТИКА, ЗВІТИ, ЗВІРКА -----------------
def history_since(hours=None, days=None):
    delta = timedelta(hours=hours) if hours else timedelta(days=days or 1)
    since = datetime.now(KYIV).replace(tzinfo=None) - delta
    return [h for h in read_history() if h["ts"] >= since]


def sane(h):
    return h["old"] <= SANE_QTY and h["new"] <= SANE_QTY


def net_of(items):
    items = [h for h in items if sane(h)]
    plus = sum(max(0, h["new"] - h["old"]) for h in items)
    minus = sum(max(0, h["old"] - h["new"]) for h in items)
    return plus, minus


def build_stats():
    df = load_df()
    total_positions = len(df)
    total_qty = int(df["q"].sum())
    zero_qty = int((df["q"] == 0).sum())

    def get_season(name):
        n = str(name).lower()
        if any(w in n for w in ["всесезон", "all weather", "all-season", "4seasons", "в.с", "в/с"]):
            return "🌦 Всесезонні"
        if any(w in n for w in ["зима", "зимов", "winter", "шип", "ice", "snow"]):
            return "❄️ Зимові"
        if any(w in n for w in ["літо", "літній", "summer"]):
            return "☀️ Літні"
        return "📦 Інші"

    def get_radius(size_str):
        m = re.search(r"R\d+", str(size_str), re.IGNORECASE)
        return m.group(0).upper() if m else "Інші"

    brands = [
        "Premiorri", "Premiori", "Rosava", "Bridgestone", "Michelin", "Nokian", "Hankook", "Continental", "Goodyear",
        "Matador", "Barum", "Cordiant", "Tigar", "Kormoran", "Sailun", "Triangle", "Windforce", "Imperial",
        "Orium", "Falken", "Dunlop", "Pirelli", "Yokohama", "Kumho", "Nexen", "BFGoodrich", "Firestone",
        "Toyo", "Lassa", "Mastersteel", "Roadstone", "Aosen", "Estrada", "Powertrac", "Petlas", "Horizon",
        "Hifly", "Marshal",
    ]

    def get_brand(name):
        n = str(name)
        for b in brands:
            if b.lower() in n.lower():
                return b
        return "Інший"

    df["season"] = df["name"].apply(get_season)
    df["radius"] = df["size"].apply(get_radius)
    df["brand"] = df["name"].apply(get_brand)
    df["loc"] = df["pallet"].apply(loc_label)
    e = html.escape

    text = (
        "📊 <b>Статистика складу шин</b>\n\n"
        f"📦 <b>Разом:</b> {total_qty} шт. ({df['code'].nunique()} артикулів, {total_positions} рядків)\n"
        f"⚠️ <b>Рядків з нульовим залишком:</b> {zero_qty}\n"
    )
    day = history_since(hours=24)
    if day:
        plus, minus = net_of(day)
        text += f"🕒 <b>За добу:</b> {len(day)} змін, +{plus} / −{minus} шт.\n"
    text += "\n❄️ <b>За сезонністю:</b>\n"
    for k, q in df.groupby("season")["q"].sum().items():
        if q > 0:
            text += f"• {k}: {int(q)} шт.\n"
    text += "\n⬡ <b>За радіусом:</b>\n"
    for k, q in sorted(df.groupby("radius")["q"].sum().items()):
        if q > 0:
            text += f"• {e(k)}: {int(q)} шт.\n"
    text += "\n🏭 <b>За виробниками:</b>\n"
    for k, q in sorted(df.groupby("brand")["q"].sum().items(), key=lambda x: x[1], reverse=True):
        if q > 0:
            text += f"• {e(k)}: {int(q)} шт.\n"
    text += "\n📍 <b>Найбільші локації:</b>\n"
    for k, q in sorted(df.groupby("loc")["q"].sum().items(), key=lambda x: x[1], reverse=True)[:6]:
        text += f"• {e(k)}: {int(q)} шт.\n"
    combo = df.groupby(["brand", "radius"])["q"].sum().reset_index()
    combo = combo[combo["q"] > 0].sort_values(by="q", ascending=False)
    text += "\n📌 <b>Топ (виробник + радіус):</b>\n"
    for _, r in combo.head(8).iterrows():
        text += f"• {e(r['brand'])} ({e(r['radius'])}): {int(r['q'])} шт.\n"
    return text


def compare_sklad(df, sk):
    tot = df.groupby("code")["q"].sum().to_dict()
    names = df.drop_duplicates("code").set_index("code")["name"].to_dict()
    out = []
    for c in set(tot) | set(sk):
        t, s = int(tot.get(c, 0)), int(sk.get(c, {}).get("qty", 0))
        if c not in sk and t == 0:
            continue
        rows = code_rows(df, c)
        where = "; ".join(f"{loc_label(r['pallet'])}: {int(r['q'])}" for _, r in rows.iterrows() if int(r["q"]) > 0) or "—"
        out.append({"code": c, "name": short_name(sk.get(c, {}).get("name") or names.get(c, "")),
                    "sklad": s, "table": t, "diff": t - s, "where": where})
    out.sort(key=lambda x: (x["diff"] == 0, -abs(x["diff"]), x["code"]))
    return out


def build_zvirka_text(rows):
    ok = [r for r in rows if r["diff"] == 0]
    less = [r for r in rows if r["diff"] < 0 and r["table"] > 0]
    more = [r for r in rows if r["diff"] > 0 and r["sklad"] > 0]
    only_sk = [r for r in rows if r["table"] == 0 and r["sklad"] != 0]
    only_tab = [r for r in rows if r["sklad"] == 0 and r["table"] > 0]
    s_sum, t_sum = sum(r["sklad"] for r in rows), sum(r["table"] for r in rows)
    text = (
        "⚖️ Звірка: склад і таблиця локацій\n\n"
        f"🏬 Склад: {s_sum} шт.   📍 Таблиця: {t_sum} шт.   Різниця: {t_sum - s_sum:+d}\n"
        f"✅ Збігається: {len(ok)} з {len(rows)} артикулів\n"
        f"🔻 У таблиці менше: {len(less)}   🔺 більше: {len(more)}\n"
        f"❓ Є на складі, немає в таблиці: {len(only_sk)}   ➕ Є в таблиці, немає на складі: {len(only_tab)}\n"
    )
    bad = [r for r in rows if r["diff"] != 0]
    if bad:
        text += "\nНайбільші розбіжності:\n"
        for r in bad[:12]:
            text += f"• {r['code']} · {r['name'][:34]} — склад {r['sklad']}, таблиця {r['table']} ({r['diff']:+d})\n"
        if len(bad) > 12:
            text += f"… і ще {len(bad) - 12}. Повний список — в Excel.\n"
    return text


def build_report_text(days=7):
    items = history_since(days=days)
    df = load_df()
    names = df.drop_duplicates("code").set_index("code")["name"].to_dict()
    text = f"📋 Звіт за {days} дн.\n\n"
    if not items:
        text += "За цей період змін не було.\n"
    else:
        plus, minus = net_of(items)
        text += f"✏️ Змін: {len(items)}   +{plus} / −{minus} шт.\n\n👥 Хто що робив:\n"
        by_user = {}
        for h in items:
            if not sane(h):
                continue
            u = by_user.setdefault(h["user"], [0, 0, 0])
            u[0] += 1
            u[1] += max(0, h["new"] - h["old"])
            u[2] += max(0, h["old"] - h["new"])
        for u, (n, p, m) in sorted(by_user.items(), key=lambda x: -x[1][0]):
            text += f"• {u}: {n} змін (+{p} / −{m})\n"
        cnt = {}
        for h in items:
            cnt[h["code"]] = cnt.get(h["code"], 0) + 1
        text += "\n🔥 Найактивніші артикули:\n"
        for c, n in sorted(cnt.items(), key=lambda x: -x[1])[:5]:
            text += f"• {c} · {short_name(names.get(c, ''))[:32]} — {n} змін\n"
    weird = [h for h in items if not sane(h)]
    if weird:
        text += f"\n⚠️ Підозрілих записів: {len(weird)} (схоже, у кількість вписали артикул), у суми не входять\n"
    moved = {h["code"] for h in items}
    stale = [c for c in df[df["q"] > 0]["code"].unique() if c not in moved]
    text += f"\n💤 Без рухів за період: {len(stale)} артикулів\n"
    zero = df[df["q"] == 0]
    text += f"⚠️ Нульові залишки: {len(zero)} рядків\n"
    return text


def build_history_text(n=10):
    items = read_history()[-n:]
    if not items:
        return "Історія змін порожня."
    lines = [f"🕘 Останні {len(items)} змін:\n"]
    for h in reversed(items):
        lines.append(f"{h['ts'].strftime('%d.%m %H:%M')} · {h['user']} · {h['code']} · {h['action']} · {h['old']} → {h['new']}")
    return "\n".join(lines)


def build_daily_report():
    df = load_df(force=True)
    text = f"🌅 Ранковий звіт, {datetime.now(KYIV).strftime('%d.%m.%Y')}\n\n"
    text += f"📦 У таблиці: {int(df['q'].sum())} шт., {df['code'].nunique()} артикулів\n"
    day = history_since(hours=24)
    if day:
        plus, minus = net_of(day)
        text += f"🕒 За добу: {len(day)} змін, +{plus} / −{minus} шт.\n"
    else:
        text += "🕒 За добу змін не було\n"
    totals = df.groupby("code")["q"].sum()
    names = df.drop_duplicates("code").set_index("code")["name"].to_dict()
    low = [(c, int(q)) for c, q in totals.items() if q <= LOW_STOCK]
    if low:
        text += f"\n🔻 Мінімальний залишок (≤ {LOW_STOCK}): {len(low)} артикулів\n"
        for c, q in sorted(low, key=lambda x: x[1])[:12]:
            text += f"• {c} · {short_name(names.get(c, ''))[:34]} — {q} шт.\n"
    sk = load_sklad_safe(force=True)
    if sk is not None:
        rows = compare_sklad(df, sk)
        bad = [r for r in rows if r["diff"] != 0]
        text += f"\n⚖️ Розбіжностей зі складом: {len(bad)} артикулів\n"
        for r in bad[:6]:
            text += f"• {r['code']} · {r['name'][:30]} — {r['diff']:+d}\n"
    return text


# ----------------- EXCEL (окремий файл, таблицю не змінює) -----------------
def _xlsx_styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style="thin", color="BFBFBF")
    return {
        "hdr_font": Font(name="Arial", bold=True, color="FFFFFF", size=11),
        "hdr_fill": PatternFill("solid", fgColor="1F3864"),
        "bad_fill": PatternFill("solid", fgColor="FCE4D6"),
        "body": Font(name="Arial", size=10),
        "bold": Font(name="Arial", size=10, bold=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _write_sheet(ws, headers, rows, widths, st, bad_rows=()):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill, c.border, c.alignment = st["hdr_font"], st["hdr_fill"], st["border"], st["center"]
    for r_i, row in enumerate(rows, 2):
        for c_i, v in enumerate(row, 1):
            c = ws.cell(r_i, c_i, v)
            c.font, c.border, c.alignment = st["body"], st["border"], st["wrap"]
            if (r_i - 2) in bad_rows:
                c.fill = st["bad_fill"]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}{max(2, len(rows) + 1)}"


def table_xlsx(df):
    from openpyxl import Workbook
    st = _xlsx_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Шини"
    rows = [[r["pallet"], int(r["code"]), r["size"], r["name"], int(r["q"])] for _, r in df.sort_values("row").iterrows()]
    _write_sheet(ws, ["№ палети / локація", "Код товару", "Розмір", "Модель", "Кількість, шт."], rows, [20, 14, 13, 56, 14], st)
    ws2 = wb.create_sheet("По артикулах")
    grouped = []
    for c in df["code"].unique():
        rr = code_rows(df, c)
        where = "; ".join(f"{loc_label(r['pallet'])}: {int(r['q'])}" for _, r in rr.iterrows() if int(r["q"]) > 0) or "—"
        grouped.append([int(c), rr.iloc[0]["size"], short_name(rr.iloc[0]["name"]), int(rr["q"].sum()), where])
    grouped.sort(key=lambda x: (-x[3], x[0]))
    _write_sheet(ws2, ["Код товару", "Розмір", "Модель", "Разом, шт.", "Де є"], grouped, [14, 13, 52, 12, 60], st)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def zvirka_xlsx(rows):
    from openpyxl import Workbook
    st = _xlsx_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Звірка"
    data = [[int(r["code"]), r["name"], r["sklad"], r["table"], None, r["where"]] for r in rows]
    _write_sheet(ws, ["Код товару", "Модель", "Склад, шт.", "Таблиця, шт.", "Різниця (таблиця − склад)", "Де є в таблиці"],
                 data, [13, 54, 11, 13, 16, 60], st, bad_rows={i for i, r in enumerate(rows) if r["diff"] != 0})
    last = len(rows) + 1
    for i in range(2, last + 1):
        ws.cell(i, 5, f"=D{i}-C{i}")
    ws.cell(last + 1, 2, "РАЗОМ").font = st["bold"]
    for col, L in ((3, "C"), (4, "D"), (5, "E")):
        ws.cell(last + 1, col, f"=SUM({L}2:{L}{last})").font = st["bold"]
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


# ----------------- ОФОРМЛЕННЯ: МЕНЮ І ДОВІДКА -----------------
MENU = ReplyKeyboardMarkup(
    [["🔍 Пошук", "➕ Додати"], ["📦 Палети", "📊 Статистика"], ["⚖️ Звірка", "📋 Звіт"], ["📥 Експорт", "❓ Допомога"]],
    resize_keyboard=True,
)

HELP_TEXT = (
    "👋 Складський помічник для обліку шин 🛞\n\n"
    "🔍 Надішли артикул, розмір (205/55R16 або «205 55 16»), назву чи бренд — покажу картку шини з усіма локаціями. "
    "Можна й фото цінника.\n"
    "➖ / ➕ під локацією змінюють кількість на 1, натискання на назву локації дозволяє ввести точну кількість.\n\n"
    "Команди:\n"
    "➕ /add — додати шину (за артикулом; розмір і назву підтягну сам)\n"
    "📦 /pallet 27 — що лежить на палеті (або /pallet зал)\n"
    "🧮 /inventory 27 — інвентаризація палети по позиціях\n"
    "📊 /stats — статистика\n"
    "⚖️ /zvirka — звірка з аркушем «Склад» (онови його й повтори команду)\n"
    "📋 /report 7 — звіт за N днів, 🕘 /history — останні зміни\n"
    "📥 /export — вивантажити таблицю в Excel\n"
    "🆔 /myid — мій ID для щоденного звіту\n"
    "❌ /cancel — скасувати поточну дію"
)


def guarded(fn):
    async def wrapper(update, context):
        try:
            return await fn(update, context)
        except Exception as e:
            msg = friendly_error(e)
            try:
                if update.callback_query:
                    await update.callback_query.answer(msg[:190], show_alert=True)
                elif update.effective_message:
                    await update.effective_message.reply_text(msg)
            except Exception:
                pass
            return None
    wrapper.__name__ = fn.__name__
    return wrapper


async def safe_edit(query, text, kb):
    try:
        await query.edit_message_text(text, reply_markup=kb)
    except Exception as e:
        if "not modified" not in str(e).lower():
            print(f"⚠️ edit: {e}")


async def notify_chats(context, text):
    for cid in REPORT_CHAT_IDS:
        try:
            await context.bot.send_message(cid, text)
        except Exception as e:
            print(f"⚠️ Не вдалося надіслати в чат {cid}: {e}")


async def after_change(context, user, row, code, old, new):
    last = context.application.bot_data.setdefault("last", {})
    last[row] = {"code": code, "old": old, "new": new, "ts": _time.time(), "uid": user.id}
    df = await asyncio.to_thread(load_df)
    new_total = code_total(df, code)
    old_total = new_total - (new - old)
    if new < old and old_total > LOW_STOCK >= new_total:
        names = df[df["code"] == code]["name"]
        name = short_name(names.iloc[0]) if len(names) else ""
        await notify_chats(context, f"🔻 Мінімальний залишок: {code} · {name} — тепер {new_total} шт. (змінив {user_label(user)})")


# ----------------- КОМАНДИ -----------------
@guarded
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(HELP_TEXT, reply_markup=MENU)


@guarded
async def myid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        f"🆔 ID цього чату: {update.effective_chat.id}\n"
        "Щоб отримувати ранковий звіт і сповіщення про мінімум, додай цей ID у змінну REPORT_CHAT_IDS на Render."
    )


@guarded
async def stats_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = await asyncio.to_thread(build_stats)
    for part in chunk_text(text):
        await update.message.reply_text(part, parse_mode="HTML")


@guarded
async def report_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    days = 7
    if context.args and context.args[0].isdigit():
        days = max(1, min(90, int(context.args[0])))
    text = await asyncio.to_thread(build_report_text, days)
    for part in chunk_text(text):
        await update.message.reply_text(part)


@guarded
async def history_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    n = 10
    if context.args and context.args[0].isdigit():
        n = max(1, min(40, int(context.args[0])))
    text = await asyncio.to_thread(build_history_text, n)
    for part in chunk_text(text):
        await update.message.reply_text(part)


@guarded
async def export_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    df = await asyncio.to_thread(load_df, True)
    bio = await asyncio.to_thread(table_xlsx, df)
    name = f"шини_{datetime.now(KYIV).strftime('%d-%m-%Y')}.xlsx"
    await update.message.reply_document(document=bio, filename=name, caption="📥 Поточна таблиця (окремий файл, Google Таблицю не змінено)")


ZVIRKA_SETUP = (
    "⚖️ Не бачу даних складу.\n\n"
    "Створи в цій Google Таблиці аркуш «{sheet}» з колонками: Артикул | Назва | Кількість, "
    "встав туди склад і повтори /zvirka."
)


@guarded
async def zvirka_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    sk = await asyncio.to_thread(load_sklad, True)
    if not sk:
        await update.effective_message.reply_text(ZVIRKA_SETUP.format(sheet=SKLAD_SHEET))
        return
    df = await asyncio.to_thread(load_df, True)
    rows = await asyncio.to_thread(compare_sklad, df, sk)
    text = build_zvirka_text(rows)
    text += f"\n🗂 Аркуш «{SKLAD_SHEET}»: {_sklad_meta['rows']} рядків, {len(sk)} артикулів."
    if _sklad_meta["dups"]:
        d = _sklad_meta["dups"]
        text += f"\nℹ️ {len(d)} артикулів на аркуші повторюються, кількості підсумовано: " + ", ".join(d[:6]) + ("…" if len(d) > 6 else "")
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("📥 Excel зі звіркою", callback_data="zvx")]])
    parts = chunk_text(text)
    for i, part in enumerate(parts):
        await update.effective_message.reply_text(part, reply_markup=kb if i == len(parts) - 1 else None)


@guarded
async def zvirka_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer("Готую файл…")
    sk = await asyncio.to_thread(load_sklad, True)
    if not sk:
        await q.message.reply_text(ZVIRKA_SETUP.format(sheet=SKLAD_SHEET))
        return
    df = await asyncio.to_thread(load_df, True)
    rows = await asyncio.to_thread(compare_sklad, df, sk)
    bio = await asyncio.to_thread(zvirka_xlsx, rows)
    name = f"звірка_{datetime.now(KYIV).strftime('%d-%m-%Y')}.xlsx"
    await q.message.reply_document(document=bio, filename=name, caption="📥 Звірка складу й таблиці локацій")


async def daily_job(context: ContextTypes.DEFAULT_TYPE):
    try:
        text = await asyncio.to_thread(build_daily_report)
    except Exception as e:
        print(f"⚠️ Помилка ранкового звіту: {e}")
        return
    await notify_chats(context, text)


# ----------------- ЛОКАЦІЇ ТА ПАЛЕТИ -----------------
async def show_location(msg, context, key):
    df = await asyncio.to_thread(load_df)
    text, kb, _rows = build_location_view(df, key)
    parts = chunk_text(text)
    for i, part in enumerate(parts):
        await msg.reply_text(part, reply_markup=kb if i == len(parts) - 1 else None)


@guarded
async def pallet_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    key = resolve_loc_key(" ".join(context.args)) if context.args else None
    if not key:
        await update.message.reply_text("📦 Обери палету або локацію:", reply_markup=palette_keyboard("pal"))
        return
    await show_location(update.message, context, key)


@guarded
async def pallet_pick(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    await show_location(q.message, context, q.data.split("_", 1)[1])


@guarded
async def pallets_back(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    await q.message.reply_text("📦 Обери палету або локацію:", reply_markup=palette_keyboard("pal"))


# ----------------- ПОШУК І КАРТКИ -----------------
async def run_search(update: Update, context: ContextTypes.DEFAULT_TYPE, query: str):
    msg = update.effective_message
    df = await asyncio.to_thread(load_df)
    codes, note = search_codes(df, query)
    if not codes:
        clean = query.strip()
        code = clean if clean.isdigit() and len(clean) >= 4 else "0"
        kb = InlineKeyboardMarkup([[InlineKeyboardButton("➕ Додати як нову шину", callback_data=f"addnew_{code}")]])
        await msg.reply_text(f"❌ За запитом «{query}» нічого не знайдено.", reply_markup=kb)
        return
    if note:
        await msg.reply_text(note)
    if len(codes) <= MAX_CARDS:
        first = df[df["code"] == codes[0]]
        photo_path = get_tire_photo(first.iloc[0]["name"]) if len(first) else None
        if photo_path and os.path.exists(photo_path):
            try:
                with open(photo_path, "rb") as f:
                    await msg.reply_photo(photo=f, caption=f"📷 {short_name(first.iloc[0]['name'])}")
            except Exception as e:
                print(f"⚠️ Не вдалося надіслати фото: {e}")
        for c in codes:
            await send_card(msg.reply_text, context, c)
    else:
        extra = f" (показано {MAX_LIST})" if len(codes) > MAX_LIST else ""
        await msg.reply_text(f"🔍 Знайшов {len(codes)} артикулів{extra}. Обери потрібний або уточни запит:",
                             reply_markup=list_keyboard(df, codes))


@guarded
async def open_card(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    code = q.data.split("_", 1)[1]
    await send_card(q.message.reply_text, context, code)


@guarded
async def refresh_card(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer("Оновлено")
    code = q.data.split("_", 1)[1]
    text, kb = await card_payload(context, code, force=True)
    await safe_edit(q, text, kb)


@guarded
async def row_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    action, row_s = q.data.split("_", 1)
    row = int(row_s)
    if action == "set":
        await q.answer()
        context.user_data["waiting_for_row"] = {"row": row, "chat": q.message.chat_id, "msg": q.message.message_id}
        await q.message.reply_text("⌨️ Надішли нову точну кількість для цієї локації (тільки число):")
        return
    code, old, new = await asyncio.to_thread(change_qty, q.from_user, row, action)
    await q.answer(f"{old} → {new} шт.")
    await after_change(context, q.from_user, row, code, old, new)
    text, kb = await card_payload(context, code)
    await safe_edit(q, text, kb)


@guarded
async def undo_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    row = int(q.data.split("_", 1)[1])
    last = context.application.bot_data.get("last", {})
    entry = last.get(row)
    if not entry or _time.time() - entry["ts"] > UNDO_WINDOW:
        await q.answer("Нічого відміняти (минуло забагато часу).", show_alert=True)
        return

    def _undo():
        ws = get_worksheet()
        vals = (ws.row_values(row) + [""] * 5)[:5]
        if str(vals[1]).strip().replace(".0", "") != entry["code"]:
            invalidate_cache()
            raise StaleRow()
        cur = to_int(vals[4])
        if cur != entry["new"]:
            return None
        ws.update_cell(row, 5, entry["old"])
        cache_set_qty(row, entry["old"])
        log_to_history(q.from_user, entry["code"], "undo", cur, entry["old"])
        return cur

    cur = await asyncio.to_thread(_undo)
    if cur is None:
        await q.answer("Після цієї зміни кількість змінювали ще раз, тому відміна небезпечна.", show_alert=True)
        return
    last.pop(row, None)
    await q.answer(f"Повернув {cur} → {entry['old']}")
    text, kb = await card_payload(context, entry["code"], force=True)
    await safe_edit(q, text, kb)


MENU_PROMPT = "🔍 Надішли артикул, розмір (наприклад 205/55R16), назву чи бренд — або фото цінника."


@guarded
async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    if not text:
        return

    # кнопки меню
    if text == "🔍 Пошук":
        await update.message.reply_text(MENU_PROMPT)
        return
    if text == "📦 Палети":
        await update.message.reply_text("📦 Обери палету або локацію:", reply_markup=palette_keyboard("pal"))
        return
    if text == "📊 Статистика":
        return await stats_command(update, context)
    if text == "⚖️ Звірка":
        return await zvirka_command(update, context)
    if text == "📋 Звіт":
        context.args = []
        return await report_command(update, context)
    if text == "📥 Експорт":
        return await export_command(update, context)
    if text == "❓ Допомога":
        await update.message.reply_text(HELP_TEXT, reply_markup=MENU)
        return

    # очікуємо точну кількість
    w = context.user_data.get("waiting_for_row")
    if w:
        if re.fullmatch(r"\d+", text):
            if int(text) > MAX_QTY:
                await update.message.reply_text(f"⚠️ {text} шт. — забагато, схоже на артикул. Введи кількість (до {MAX_QTY}) ще раз:")
                return
            context.user_data.pop("waiting_for_row", None)
            code, old, new = await asyncio.to_thread(change_qty, update.effective_user, w["row"], "set", int(text))
            await after_change(context, update.effective_user, w["row"], code, old, new)
            await update.message.reply_text(f"✅ {code}: було {old}, стало {new} шт.")
            card_text, kb = await card_payload(context, code)
            try:
                await context.bot.edit_message_text(card_text, chat_id=w["chat"], message_id=w["msg"], reply_markup=kb)
            except Exception:
                await update.message.reply_text(card_text, reply_markup=kb)
            return
        context.user_data.pop("waiting_for_row", None)  # це не число — вважаємо новим пошуком

    # «палета 27», «зал», «фасад»
    low = text.lower()
    m = re.fullmatch(r"(?:палета|палет|пал\.?|№)\s*(\d+)", low)
    key = m.group(1) if m else (resolve_loc_key(low) if low in SPECIAL_LOCS.values() or low in SPECIAL_LOCS else None)
    if key:
        await show_location(update.message, context, key)
        return

    await run_search(update, context, text)


@guarded
async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not openai_client:
        await update.message.reply_text("⚠️ Пошук по фото не налаштовано. Надішли артикул текстом.")
        return
    msg = await update.message.reply_text("🔍 Аналізую цінник і шукаю артикул...")
    file_path = f"photo_{update.message.chat_id}_{update.message.message_id}.jpg"

    def ocr():
        with open(file_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
        resp = openai_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": "Знайди на цьому ціннику 8-значний артикул після слів 'арт.' або 'арт'. Поверни ТІЛЬКИ ці цифри і нічого більше. Якщо немає, поверни 0."},
                    {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}},
                ],
            }],
            max_tokens=20,
        )
        return resp.choices[0].message.content.strip()

    try:
        photo_file = await update.message.photo[-1].get_file()
        await photo_file.download_to_drive(file_path)
        code = re.sub(r"\D", "", await asyncio.to_thread(ocr))
        if not code or code == "0":
            await msg.edit_text("❌ Не вдалося знайти артикул на фото.")
            return
        await msg.edit_text(f"🏷 Розпізнано артикул: {code}")
        await run_search(update, context, code)
    except Exception as e:
        await msg.edit_text(friendly_error(e))
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)


async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("🎙 Голосові повідомлення не підтримуються. Напиши текстом.")


async def cancel_plain(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("waiting_for_row", None)
    await update.message.reply_text("Скасовано.", reply_markup=MENU)


# ----------------- ДОДАВАННЯ ШИНИ (/add): артикул → палета → кількість -----------------
CODE, NEWCONFIRM, SIZE, NAME, PALLET, QTY, CONFIRM = range(7)


def parse_pallet(text):
    m = re.fullmatch(r"(?:палет[аиуі]?\s*)?№?\s*(\d+)", text.strip(), re.IGNORECASE)
    return m.group(1) if m else text.strip()


def locations_text(rows):
    return "\n".join(f"• {loc_label(r['pallet'])} — {int(r['q'])} шт." for _, r in rows.head(10).iterrows())


def find_same(df, code, pallet):
    same = df[(df["code"] == code) & (df["pallet"].str.strip().str.lower() == pallet.strip().lower())]
    return [int(x) for x in same["row"]]


@guarded
async def add_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("waiting_for_row", None)
    context.user_data["new"] = {}
    if update.callback_query:
        q = update.callback_query
        await q.answer()
        code = q.data.split("_", 1)[1]
        if code.isdigit() and code != "0":
            context.user_data["new"]["code"] = code
            await q.message.reply_text(f"➕ Нова шина, артикул {code} (/cancel — скасувати).\n📏 Розмір? (наприклад 205/55 R16)")
            return SIZE
        await q.message.reply_text("➕ Додаємо шину (/cancel — скасувати).\n🔢 Артикул?")
        return CODE
    await update.message.reply_text("➕ Додаємо шину (/cancel — скасувати).\n🔢 Артикул?")
    return CODE


@guarded
async def add_code(update: Update, context: ContextTypes.DEFAULT_TYPE):
    code = re.sub(r"\D", "", update.message.text or "")
    if len(code) < 4:
        await update.message.reply_text("⚠️ Артикул має містити цифри (мінімум 4). Спробуй ще раз:")
        return CODE
    new = context.user_data.setdefault("new", {})
    new["code"] = code
    df = await asyncio.to_thread(load_df)
    rows = code_rows(df, code)
    if not rows.empty:
        first = rows.iloc[0]
        new["size"], new["name"] = first["size"], first["name"]
        await update.message.reply_text(
            "✅ Знайдено в таблиці:\n"
            f"• Артикул: {code}\n• Розмір: {first['size']}\n• Назва: {short_name(first['name'])}\n\n"
            f"Вже є:\n{locations_text(rows)}\n\n📍 На яку палету додаємо?"
        )
        return PALLET
    kb = InlineKeyboardMarkup([[
        InlineKeyboardButton("➕ Додати як нову", callback_data="nt_yes"),
        InlineKeyboardButton("❌ Скасувати", callback_data="nt_no"),
    ]])
    await update.message.reply_text(f"❌ Артикула {code} в таблиці немає. Додати як нову шину?", reply_markup=kb)
    return NEWCONFIRM


@guarded
async def add_new_decision(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    new = context.user_data.get("new")
    if q.data == "nt_no" or not new:
        context.user_data.pop("new", None)
        await q.edit_message_text("❌ Скасовано.")
        return ConversationHandler.END
    await q.edit_message_text(f"➕ Нова шина, артикул {new['code']}")
    await q.message.reply_text("📏 Розмір? (наприклад 205/55 R16)")
    return SIZE


@guarded
async def add_size(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.setdefault("new", {})["size"] = (update.message.text or "").strip()
    await update.message.reply_text("🏷 Назва (бренд, модель, сезон)?")
    return NAME


@guarded
async def add_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = (update.message.text or "").strip()
    if not name.lower().startswith("автошина"):
        name = f"Автошина {name}"
    context.user_data.setdefault("new", {})["name"] = name
    await update.message.reply_text("📍 На яку палету додаємо?")
    return PALLET


@guarded
async def add_pallet(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    if not text:
        await update.message.reply_text("⚠️ Вкажи палету або локацію:")
        return PALLET
    context.user_data.setdefault("new", {})["pallet"] = parse_pallet(text)
    await update.message.reply_text("🔢 Кількість (шт.)?")
    return QTY


@guarded
async def add_qty(update: Update, context: ContextTypes.DEFAULT_TYPE):
    m = re.fullmatch(r"\s*(\d+)\s*", update.message.text or "")
    if not m or int(m.group(1)) > MAX_QTY:
        await update.message.reply_text(f"⚠️ Введи кількість числом (до {MAX_QTY}):")
        return QTY
    new = context.user_data.setdefault("new", {})
    new["qty"] = int(m.group(1))
    if not all(k in new for k in ("pallet", "code", "size", "name")):
        await update.message.reply_text("⚠️ Щось пішло не так. Почни спочатку: /add")
        context.user_data.pop("new", None)
        return ConversationHandler.END

    df = await asyncio.to_thread(load_df)
    sk = await asyncio.to_thread(load_sklad_safe)
    dup = find_same(df, new["code"], new["pallet"])

    summary = (
        "Перевір і підтверди:\n\n"
        f"📍 Палета: {new['pallet']}\n"
        f"• Артикул: {new['code']}\n"
        f"• Розмір: {new['size']}\n"
        f"• Назва: {short_name(new['name'])}\n"
        f"• Кількість: {new['qty']} шт."
    )
    if sk is not None and new["code"] not in sk:
        summary += "\n\n⚠️ Такого артикула немає в базі складу. Переконайся, що код правильний."
    if dup:
        new["merge_row"] = dup[0]
        summary += (f"\n\n⚠️ Цей артикул на палеті {new['pallet']} уже є (рядок {dup[0]}). "
                    "Додати кількість до наявного рядка чи створити окремий?")
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton(f"➕ Додати +{new['qty']} до наявного", callback_data="add_merge")],
            [InlineKeyboardButton("📄 Окремий новий рядок", callback_data="add_ok")],
            [InlineKeyboardButton("❌ Скасувати", callback_data="add_no")],
        ])
    else:
        kb = InlineKeyboardMarkup([[
            InlineKeyboardButton("✅ Додати", callback_data="add_ok"),
            InlineKeyboardButton("❌ Скасувати", callback_data="add_no"),
        ]])
    await update.message.reply_text(summary, reply_markup=kb)
    return CONFIRM


@guarded
async def add_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    new = context.user_data.pop("new", None)
    if q.data == "add_no" or not new:
        await q.edit_message_text("❌ Скасовано.")
        return ConversationHandler.END
    try:
        if q.data == "add_merge" and new.get("merge_row"):
            code, old, total = await asyncio.to_thread(merge_qty, q.from_user, new["merge_row"], new["qty"])
            await q.edit_message_text(
                f"✅ Додано {new['qty']} шт. до рядка {new['merge_row']} (палета {new['pallet']}, арт. {code}): "
                f"було {old}, стало {total}."
            )
        else:
            row = await asyncio.to_thread(add_tire, q.from_user, new["pallet"], new["code"], new["size"], new["name"], new["qty"])
            await q.edit_message_text(
                f"✅ Додано в рядок {row}: {short_name(new['name'])} ({new['size']}), арт. {new['code']}, "
                f"палета {new['pallet']}, {new['qty']} шт."
            )
        await send_card(q.message.reply_text, context, new["code"], force=True)
    except Exception as e:
        await q.edit_message_text(friendly_error(e))
    return ConversationHandler.END


async def add_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("new", None)
    await update.message.reply_text("❌ Додавання скасовано.", reply_markup=MENU)
    return ConversationHandler.END


async def add_timeout(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("new", None)
    return ConversationHandler.END


# ----------------- ІНВЕНТАРИЗАЦІЯ ПАЛЕТИ (/inventory) -----------------
INV_ASK, INV_CONFIRM = 20, 21


def inv_keyboard(old):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(f"✅ Так, {old} шт.", callback_data="ivok")],
        [InlineKeyboardButton("⏭ Пропустити", callback_data="ivskip"), InlineKeyboardButton("🛑 Завершити", callback_data="ivend")],
    ])


async def inv_strip_buttons(context, inv):
    if inv.get("msg"):
        try:
            await context.bot.edit_message_reply_markup(chat_id=inv["chat"], message_id=inv["msg"], reply_markup=None)
        except Exception:
            pass


async def inv_ask(context, inv):
    items, idx = inv["items"], inv["idx"]
    if idx >= len(items):
        return await inv_finish(context, inv)
    it = items[idx]
    text = (
        f"🧮 Інвентаризація: {inv['label']} ({idx + 1}/{len(items)})\n\n"
        f"🛞 {it['name']}\n🔢 {it['code']}   📏 {it['size']}\n"
        f"📋 У таблиці: {it['old']} шт.\n\n"
        "Скільки лежить насправді? Напиши число або натисни кнопку."
    )
    sent = await context.bot.send_message(inv["chat"], text, reply_markup=inv_keyboard(it["old"]))
    inv["msg"] = sent.message_id
    return INV_ASK


async def inv_finish(context, inv):
    diffs = [(it, inv["res"][it["row"]]) for it in inv["items"] if it["row"] in inv["res"] and inv["res"][it["row"]] != it["old"]]
    checked = len(inv["res"])
    if not diffs:
        await context.bot.send_message(inv["chat"], f"✅ Інвентаризація {inv['label']}: перевірено {checked} позицій, розбіжностей немає.", reply_markup=MENU)
        inv["done"] = True
        return ConversationHandler.END
    lines = [f"🧮 Підсумок інвентаризації: {inv['label']}\nПеревірено: {checked} з {len(inv['items'])} позицій.\n\nРозбіжності:"]
    for it, new in diffs:
        lines.append(f"• {it['code']} · {it['size']}: {it['old']} → {new} ({new - it['old']:+d})")
    lines.append("\nЗастосувати ці зміни в таблиці?")
    kb = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ Застосувати", callback_data="ivapply"),
        InlineKeyboardButton("❌ Скасувати", callback_data="ivcancel"),
    ]])
    inv["diffs"] = [(it["row"], it["code"], it["old"], new) for it, new in diffs]
    await context.bot.send_message(inv["chat"], "\n".join(lines), reply_markup=kb)
    return INV_CONFIRM


async def inv_begin(update: Update, context: ContextTypes.DEFAULT_TYPE, key):
    context.user_data.pop("waiting_for_row", None)
    chat_id = update.effective_chat.id
    if update.callback_query:
        await update.callback_query.answer()
    df = await asyncio.to_thread(load_df, True)
    rows = location_rows(df, key)
    if rows.empty:
        await context.bot.send_message(chat_id, f"📦 {loc_name_of_key(key)}: нічого інвентаризувати.")
        return ConversationHandler.END
    items = [{"row": int(r["row"]), "code": r["code"], "size": r["size"], "name": short_name(r["name"]), "old": int(r["q"])}
             for _, r in rows.iterrows()]
    inv = {"key": key, "label": loc_name_of_key(key), "items": items, "idx": 0, "res": {}, "chat": chat_id, "msg": None}
    context.user_data["inv"] = inv
    await context.bot.send_message(chat_id, f"🧮 Починаємо інвентаризацію: {inv['label']}, {len(items)} позицій. /cancel — перервати без змін.")
    return await inv_ask(context, inv)


@guarded
async def inv_entry_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    key = resolve_loc_key(" ".join(context.args)) if context.args else None
    if not key:
        await update.message.reply_text("🧮 Інвентаризація: обери палету або локацію:", reply_markup=palette_keyboard("invs"))
        return ConversationHandler.END
    return await inv_begin(update, context, key)


@guarded
async def inv_entry_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    return await inv_begin(update, context, update.callback_query.data.split("_", 1)[1])


async def _inv_current(context):
    inv = context.user_data.get("inv")
    if not inv or inv["idx"] >= len(inv["items"]):
        return None, None
    return inv, inv["items"][inv["idx"]]


@guarded
async def inv_button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    inv, it = await _inv_current(context)
    if not inv:
        return ConversationHandler.END
    await inv_strip_buttons(context, inv)
    if q.data == "ivend":
        return await inv_finish(context, inv)
    if q.data == "ivok":
        inv["res"][it["row"]] = it["old"]
    inv["idx"] += 1
    return await inv_ask(context, inv)


@guarded
async def inv_number(update: Update, context: ContextTypes.DEFAULT_TYPE):
    inv, it = await _inv_current(context)
    if not inv:
        return ConversationHandler.END
    m = re.fullmatch(r"\s*(\d+)\s*", update.message.text or "")
    if not m or int(m.group(1)) > MAX_QTY:
        await update.message.reply_text(f"⚠️ Напиши кількість числом (до {MAX_QTY}) або натисни кнопку.")
        return INV_ASK
    await inv_strip_buttons(context, inv)
    inv["res"][it["row"]] = int(m.group(1))
    inv["idx"] += 1
    return await inv_ask(context, inv)


@guarded
async def inv_confirm(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    inv = context.user_data.pop("inv", None)
    if q.data == "ivcancel" or not inv:
        await q.edit_message_text("❌ Зміни не застосовано.")
        return ConversationHandler.END
    try:
        n = await asyncio.to_thread(apply_qty_batch, q.from_user, inv["diffs"], "inv")
        await q.edit_message_text(f"✅ Оновлено {n} позицій ({inv['label']}).")
        for _row, code, old, new in inv["diffs"]:
            if new < old:
                df = await asyncio.to_thread(load_df)
                total = code_total(df, code)
                if total <= LOW_STOCK < total + (old - new):
                    await notify_chats(context, f"🔻 Мінімальний залишок: {code} — тепер {total} шт. (інвентаризація, {user_label(q.from_user)})")
    except Exception as e:
        await q.edit_message_text(friendly_error(e))
    return ConversationHandler.END


async def inv_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("inv", None)
    await update.message.reply_text("❌ Інвентаризацію перервано, змін у таблиці немає.", reply_markup=MENU)
    return ConversationHandler.END


async def inv_timeout(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("inv", None)
    return ConversationHandler.END


async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE):
    print(f"⚠️ Сталася помилка при обробці оновлення: {context.error}")


async def post_init(app):
    try:
        await app.bot.set_my_commands([
            ("start", "Меню й довідка"),
            ("add", "Додати шину"),
            ("pallet", "Що лежить на палеті"),
            ("inventory", "Інвентаризація палети"),
            ("stats", "Статистика"),
            ("zvirka", "Звірка зі складом"),
            ("report", "Звіт за N днів"),
            ("history", "Останні зміни"),
            ("export", "Вивантажити в Excel"),
            ("myid", "Мій ID для звітів"),
            ("cancel", "Скасувати дію"),
        ])
    except Exception as e:
        print(f"⚠️ Не вдалося встановити команди: {e}")


# ----------------- ЗАПУСК -----------------
def main():
    if TELEGRAM_BOT_TOKEN == "ВСТАВ_СЮДИ_ТОКЕН":
        print("❌ Не задано токен: додай змінну TELEGRAM_BOT_TOKEN або встав токен у bot.py")
        sys.exit(1)

    app = ApplicationBuilder().token(TELEGRAM_BOT_TOKEN).post_init(post_init).build()
    plain_text = filters.TEXT & ~filters.COMMAND
    MENU_RE = r"^(🔍 Пошук|➕ Додати|📦 Палети|📊 Статистика|⚖️ Звірка|📋 Звіт|📥 Експорт|❓ Допомога)$"
    conv_text = plain_text & ~filters.Regex(MENU_RE)

    add_conv = ConversationHandler(
        entry_points=[
            CommandHandler("add", add_start),
            MessageHandler(filters.Regex(r"^➕ Додати$"), add_start),
            CallbackQueryHandler(add_start, pattern=r"^addnew_\d+$"),
        ],
        states={
            CODE: [MessageHandler(conv_text, add_code)],
            NEWCONFIRM: [CallbackQueryHandler(add_new_decision, pattern=r"^nt_(yes|no)$")],
            SIZE: [MessageHandler(conv_text, add_size)],
            NAME: [MessageHandler(conv_text, add_name)],
            PALLET: [MessageHandler(conv_text, add_pallet)],
            QTY: [MessageHandler(conv_text, add_qty)],
            CONFIRM: [CallbackQueryHandler(add_confirm, pattern=r"^add_(ok|no|merge)$")],
            ConversationHandler.TIMEOUT: [MessageHandler(filters.ALL, add_timeout)],
        },
        fallbacks=[CommandHandler("cancel", add_cancel)],
        conversation_timeout=900,
        name="add_conv",
        allow_reentry=True,
    )

    inv_conv = ConversationHandler(
        entry_points=[
            CommandHandler("inventory", inv_entry_cmd),
            CallbackQueryHandler(inv_entry_cb, pattern=r"^invs_(\d+|zal|vul|fas)$"),
        ],
        states={
            INV_ASK: [
                CallbackQueryHandler(inv_button, pattern=r"^iv(ok|skip|end)$"),
                MessageHandler(conv_text, inv_number),
            ],
            INV_CONFIRM: [CallbackQueryHandler(inv_confirm, pattern=r"^iv(apply|cancel)$")],
            ConversationHandler.TIMEOUT: [MessageHandler(filters.ALL, inv_timeout)],
        },
        fallbacks=[CommandHandler("cancel", inv_cancel)],
        conversation_timeout=1800,
        name="inv_conv",
    )

    app.add_handler(add_conv)   # розмови — першими
    app.add_handler(inv_conv)
    app.add_handler(CommandHandler(["start", "help"], start))
    app.add_handler(CommandHandler("stats", stats_command))
    app.add_handler(CommandHandler("pallet", pallet_command))
    app.add_handler(CommandHandler("zvirka", zvirka_command))
    app.add_handler(CommandHandler("report", report_command))
    app.add_handler(CommandHandler("history", history_command))
    app.add_handler(CommandHandler("export", export_command))
    app.add_handler(CommandHandler("myid", myid))
    app.add_handler(CommandHandler("cancel", cancel_plain))
    app.add_handler(CallbackQueryHandler(row_action, pattern=r"^(inc|dec|set)_\d+$"))
    app.add_handler(CallbackQueryHandler(undo_action, pattern=r"^undo_\d+$"))
    app.add_handler(CallbackQueryHandler(refresh_card, pattern=r"^rf_\d+$"))
    app.add_handler(CallbackQueryHandler(open_card, pattern=r"^card_\d+$"))
    app.add_handler(CallbackQueryHandler(pallet_pick, pattern=r"^pal_(\d+|zal|vul|fas)$"))
    app.add_handler(CallbackQueryHandler(pallets_back, pattern=r"^pals$"))
    app.add_handler(CallbackQueryHandler(zvirka_excel, pattern=r"^zvx$"))
    app.add_handler(MessageHandler(plain_text, handle_message))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.VOICE, handle_voice))
    app.add_error_handler(error_handler)

    if REPORT_CHAT_IDS:
        if app.job_queue is None:
            print("⚠️ JobQueue недоступний: додай python-telegram-bot[job-queue] у requirements.txt для ранкового звіту")
        else:
            try:
                hh, mm = [int(x) for x in REPORT_TIME.split(":")]
            except Exception:
                hh, mm = 8, 30
            app.job_queue.run_daily(daily_job, time=dtime(hour=hh, minute=mm, tzinfo=KYIV), name="daily_report")
            print(f"🕗 Ранковий звіт щодня о {hh:02d}:{mm:02d} (Київ) → {REPORT_CHAT_IDS}")

    print("🤖 Бот успішно запущений!")
    app.run_polling()


if __name__ == "__main__":
    main()
